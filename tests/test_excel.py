import pytest
from openpyxl import load_workbook

from modules.documents import excel


def use_temp_documents(tmp_path, monkeypatch):
    monkeypatch.setattr(excel, "DOCUMENT_ROOT", tmp_path)


def test_safe_path_rejects_traversal(tmp_path, monkeypatch):
    use_temp_documents(tmp_path, monkeypatch)

    with pytest.raises(ValueError):
        excel._safe_document_path("../outside.xlsx")

    with pytest.raises(ValueError):
        excel._safe_document_path("..\\outside.xlsx")

    with pytest.raises(ValueError):
        excel._safe_document_path("C:\\outside.xlsx")


def test_create_and_read_excel(tmp_path, monkeypatch):
    use_temp_documents(tmp_path, monkeypatch)

    created = excel.excel_create(
        "rapport.xlsx",
        sheet_name="Data",
        headers=["Namn", "Värde"],
    )
    text = excel.excel_read(
        "rapport.xlsx",
        sheet_name="Data",
    )

    assert "rapport.xlsx" in created
    assert "blad: Data" in text
    assert "Namn\tVärde" in text


def test_create_refuses_overwrite_by_default(tmp_path, monkeypatch):
    use_temp_documents(tmp_path, monkeypatch)
    excel.excel_create("rapport.xlsx")

    with pytest.raises(FileExistsError):
        excel.excel_create("rapport.xlsx")


def test_write_cell_updates_existing_workbook(tmp_path, monkeypatch):
    use_temp_documents(tmp_path, monkeypatch)
    excel.excel_create("rapport.xlsx", sheet_name="Data")

    excel.excel_write_cell(
        "rapport.xlsx",
        "B2",
        42,
        sheet_name="Data",
    )

    workbook = load_workbook(tmp_path / "rapport.xlsx")
    try:
        assert workbook["Data"]["B2"].value == 42
    finally:
        workbook.close()


def test_append_row_adds_values(tmp_path, monkeypatch):
    use_temp_documents(tmp_path, monkeypatch)
    excel.excel_create(
        "rapport.xlsx",
        headers=["A", "B"],
    )

    excel.excel_append_row(
        "rapport.xlsx",
        ["x", 7],
    )

    workbook = load_workbook(tmp_path / "rapport.xlsx")
    try:
        sheet = workbook.active
        assert sheet["A2"].value == "x"
        assert sheet["B2"].value == 7
    finally:
        workbook.close()


def test_formulas_are_blocked_by_default(tmp_path, monkeypatch):
    use_temp_documents(tmp_path, monkeypatch)
    excel.excel_create("rapport.xlsx")

    with pytest.raises(ValueError):
        excel.excel_write_cell(
            "rapport.xlsx",
            "A1",
            "=SUM(B1:B2)",
        )


def test_formula_can_be_explicitly_allowed(tmp_path, monkeypatch):
    use_temp_documents(tmp_path, monkeypatch)
    excel.excel_create("rapport.xlsx")

    excel.excel_write_cell(
        "rapport.xlsx",
        "A1",
        "=SUM(B1:B2)",
        allow_formula=True,
    )

    workbook = load_workbook(
        tmp_path / "rapport.xlsx",
        data_only=False,
    )
    try:
        assert workbook.active["A1"].value == "=SUM(B1:B2)"
    finally:
        workbook.close()


def test_sheet_name_validation():
    with pytest.raises(ValueError):
        excel._validate_sheet_name("bad/name")


def test_excel_tools_declare_parameter_schemas():
    for name in (
        "excel_create",
        "excel_read",
        "excel_write_cell",
        "excel_append_row",
    ):
        assert "parameters" in excel.TOOLS[name]
