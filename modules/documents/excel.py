from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.config import PROJECT_ROOT


DOCUMENT_ROOT = PROJECT_ROOT / "runtime" / "documents"
INVALID_SHEET_CHARS = set("[]:*?/\\")
SAFE_SCALAR_TYPES = (str, int, float, bool, type(None))


def _safe_document_path(filename):
    value = (filename or "").strip().replace("\\", "/")

    if not value:
        raise ValueError("Filnamn saknas.")

    relative = Path(value)

    if relative.is_absolute():
        raise ValueError("Absoluta sökvägar är inte tillåtna.")

    first_part = relative.parts[0] if relative.parts else ""

    if ":" in first_part:
        raise ValueError("Enhetsbeteckning i sökvägen är inte tillåten.")

    if relative.suffix.lower() != ".xlsx":
        raise ValueError("Excel-filen måste ha filändelsen .xlsx.")

    root = DOCUMENT_ROOT.resolve()
    candidate = (root / relative).resolve()

    if candidate != root and root not in candidate.parents:
        raise ValueError("Sökvägen lämnar den tillåtna dokumentmappen.")

    return candidate


def _validate_sheet_name(sheet_name):
    name = (sheet_name or "").strip()

    if not name:
        raise ValueError("Bladnamn saknas.")

    if len(name) > 31:
        raise ValueError("Excel-bladnamn får vara högst 31 tecken.")

    if any(char in INVALID_SHEET_CHARS for char in name):
        raise ValueError("Bladnamnet innehåller otillåtna tecken.")

    return name


def _validate_cell_value(value, allow_formula=False):
    if not isinstance(value, SAFE_SCALAR_TYPES):
        raise ValueError(
            "Excel-celler får endast innehålla text, tal, booleskt värde eller tomt värde."
        )

    if (
        isinstance(value, str)
        and value.startswith("=")
        and not allow_formula
    ):
        raise ValueError(
            "Formler är blockerade som standard. Sätt allow_formula=true om formeln är avsiktlig."
        )

    return value


def _sheet(workbook, sheet_name=None):
    if sheet_name is None:
        return workbook.active

    name = _validate_sheet_name(sheet_name)

    if name not in workbook.sheetnames:
        raise KeyError(f"Excel-bladet finns inte: {name}")

    return workbook[name]


def excel_create(
    filename,
    sheet_name="Sheet1",
    headers=None,
    overwrite=False,
):
    path = _safe_document_path(filename)
    name = _validate_sheet_name(sheet_name)

    if path.exists() and not overwrite:
        raise FileExistsError(
            f"Filen finns redan: {path.name}. Sätt overwrite=true för att ersätta den."
        )

    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = name

    if headers is not None:
        if not isinstance(headers, list):
            raise ValueError("headers måste vara en lista.")

        normalized = []

        for header in headers:
            if not isinstance(header, str):
                raise ValueError("Alla rubriker måste vara text.")

            normalized.append(
                _validate_cell_value(
                    header,
                    allow_formula=False,
                )
            )

        if normalized:
            sheet.append(normalized)

    workbook.save(path)
    workbook.close()

    return f"Excel-fil skapad: {path}"


def excel_read(filename, sheet_name=None, max_rows=20):
    path = _safe_document_path(filename)

    if not path.exists():
        raise FileNotFoundError(f"Excel-filen finns inte: {path.name}")

    limit = max(1, min(int(max_rows), 100))
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=False,
    )

    try:
        sheet = _sheet(workbook, sheet_name)
        lines = [
            f"Excel: {path.name} | blad: {sheet.title}"
        ]

        row_count = 0

        for row in sheet.iter_rows(values_only=True):
            row_count += 1

            if row_count > limit:
                break

            values = [
                "" if value is None else str(value)
                for value in row[:50]
            ]
            lines.append("\t".join(values))

        if row_count == 0:
            lines.append("(bladet är tomt)")

        if row_count > limit:
            lines.append(
                f"... visningen begränsades till {limit} rader."
            )

        return "\n".join(lines)
    finally:
        workbook.close()


def excel_write_cell(
    filename,
    cell,
    value,
    sheet_name=None,
    allow_formula=False,
):
    path = _safe_document_path(filename)

    if not path.exists():
        raise FileNotFoundError(f"Excel-filen finns inte: {path.name}")

    normalized_value = _validate_cell_value(
        value,
        allow_formula=allow_formula,
    )

    workbook = load_workbook(path)

    try:
        sheet = _sheet(workbook, sheet_name)
        sheet[cell] = normalized_value
        workbook.save(path)
    finally:
        workbook.close()

    return (
        f"Excel-cell uppdaterad: {path.name} | "
        f"{sheet_name or 'aktivt blad'}!{cell}"
    )


def excel_append_row(
    filename,
    values,
    sheet_name=None,
    allow_formula=False,
):
    path = _safe_document_path(filename)

    if not path.exists():
        raise FileNotFoundError(f"Excel-filen finns inte: {path.name}")

    if not isinstance(values, list):
        raise ValueError("values måste vara en lista.")

    if len(values) > 100:
        raise ValueError("En rad får innehålla högst 100 värden.")

    normalized = [
        _validate_cell_value(
            value,
            allow_formula=allow_formula,
        )
        for value in values
    ]

    workbook = load_workbook(path)

    try:
        sheet = _sheet(workbook, sheet_name)
        sheet.append(normalized)
        workbook.save(path)
        row_number = sheet.max_row
    finally:
        workbook.close()

    return (
        f"Rad tillagd i {path.name} | "
        f"{sheet_name or 'aktivt blad'} | rad {row_number}"
    )


TOOLS = {
    "excel_create": {
        "function": excel_create,
        "description": (
            "Skapar en lokal .xlsx-fil under runtime/documents. "
            "Kan lägga till rubriker och ersätter inte befintlig fil utan uttryckligt overwrite."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string"},
                "sheet_name": {"type": "string"},
                "headers": {"type": "array"},
                "overwrite": {"type": "boolean"},
            },
            "required": ["filename"],
            "additionalProperties": False,
        },
    },
    "excel_read": {
        "function": excel_read,
        "description": (
            "Läser upp till ett begränsat antal rader från en lokal .xlsx-fil under runtime/documents."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string"},
                "sheet_name": {"type": "string"},
                "max_rows": {"type": "integer"},
            },
            "required": ["filename"],
            "additionalProperties": False,
        },
    },
    "excel_write_cell": {
        "function": excel_write_cell,
        "description": (
            "Skriver ett värde till en specificerad Excel-cell i en lokal arbetsbok. "
            "Formler blockeras om allow_formula inte uttryckligen är true."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string"},
                "cell": {"type": "string"},
                "value": {},
                "sheet_name": {"type": "string"},
                "allow_formula": {"type": "boolean"},
            },
            "required": ["filename", "cell", "value"],
            "additionalProperties": False,
        },
    },
    "excel_append_row": {
        "function": excel_append_row,
        "description": (
            "Lägger till en rad i en lokal Excel-arbetsbok. "
            "Formler blockeras om allow_formula inte uttryckligen är true."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string"},
                "values": {"type": "array"},
                "sheet_name": {"type": "string"},
                "allow_formula": {"type": "boolean"},
            },
            "required": ["filename", "values"],
            "additionalProperties": False,
        },
    },
}
