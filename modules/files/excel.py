import os
import re
from pathlib import Path

from core.config import load_settings
from modules.files.workspace import (
    extract_file_path,
    resolve_workspace_path,
)


def _load_openpyxl():
    try:
        import openpyxl
    except ImportError as error:
        raise RuntimeError(
            "Excel-funktioner kräver openpyxl. "
            "Installera requirements-excel.txt."
        ) from error

    return openpyxl


def _excel_config(settings=None):
    settings = settings or load_settings()
    files = settings.get("files", {})
    excel = settings.get("excel", {})

    if not files.get("enabled", True):
        raise RuntimeError("Filverktygen är avstängda i konfigurationen.")

    if not excel.get("enabled", True):
        raise RuntimeError("Excel-verktygen är avstängda i konfigurationen.")

    return settings, excel


def _require_xlsx(path_text, settings=None):
    path = resolve_workspace_path(path_text, settings)

    if path.suffix.lower() != ".xlsx":
        raise ValueError("Excel-verktyget tillåter endast .xlsx-filer.")

    return path


def _coerce_cell(value):
    text = str(value).strip()

    if text == "":
        return ""

    if text.startswith("="):
        return text

    lowered = text.lower()

    if lowered == "true":
        return True

    if lowered == "false":
        return False

    try:
        return int(text)
    except ValueError:
        pass

    try:
        return float(text.replace(",", "."))
    except ValueError:
        return text


def _split_cells(text):
    return [
        _coerce_cell(item)
        for item in str(text).split(",")
    ]


def parse_create_request(user_text):
    text = str(user_text or "").strip()
    path = extract_file_path(text)

    if not path:
        raise ValueError("Ingen .xlsx-fil hittades i instruktionen.")

    if not path.lower().endswith(".xlsx"):
        raise ValueError("Filnamnet måste sluta med .xlsx.")

    columns_match = re.search(
        r"kolumner\s+(.+?)(?=\s+(?:och\s+)?rader\s+)",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    rows_match = re.search(
        r"(?:och\s+)?rader\s+(.+)$",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    if not columns_match:
        raise ValueError(
            "Kolumner saknas. Ange exempelvis: kolumner Namn, Belopp."
        )

    headers = [
        str(value).strip()
        for value in columns_match.group(1).split(",")
        if str(value).strip()
    ]

    if not headers:
        raise ValueError("Minst en kolumn krävs.")

    rows = []

    if rows_match:
        for row_text in rows_match.group(1).split(";"):
            if row_text.strip():
                rows.append(_split_cells(row_text))

    for row in rows:
        if len(row) != len(headers):
            raise ValueError(
                "Varje rad måste ha samma antal värden som kolumner."
            )

    return {
        "path": path,
        "headers": headers,
        "rows": rows,
    }


def parse_append_request(user_text):
    text = str(user_text or "").strip()
    path = extract_file_path(text)

    if not path or not path.lower().endswith(".xlsx"):
        raise ValueError("Ingen giltig .xlsx-fil hittades i instruktionen.")

    escaped = re.escape(path)
    match = re.search(
        rf"(?:lägg\s+till|addera)\s+(?:raden\s+)?(.+?)"
        rf"\s+i\s+(?:filen\s+)?[\"'“”]?{escaped}[\"'“”]?",
        text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    if not match:
        raise ValueError(
            'Använd exempelvis: Lägg till raden Kaffe, 35 i "budget.xlsx".'
        )

    row = _split_cells(match.group(1))

    return {
        "path": path,
        "row": row,
    }


def create_excel_file(
    request,
    settings=None,
    openpyxl_module=None,
):
    settings, excel = _excel_config(settings)

    if not excel.get("write_enabled", False):
        raise RuntimeError(
            "Excel-skrivning är avstängd i konfigurationen."
        )

    path = _require_xlsx(request["path"], settings)
    path.parent.mkdir(parents=True, exist_ok=True)
    openpyxl = openpyxl_module or _load_openpyxl()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = excel.get("default_sheet", "Data")
    sheet.append(list(request["headers"]))

    for row in request.get("rows", []):
        sheet.append(list(row))

    sheet.freeze_panes = "A2"
    temporary = path.with_suffix(".tmp.xlsx")

    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()

    return {
        "path": path,
        "sheet": sheet.title,
        "rows_written": len(request.get("rows", [])),
        "columns": len(request["headers"]),
    }


def read_excel_file(
    path_text,
    settings=None,
    openpyxl_module=None,
):
    settings, excel = _excel_config(settings)
    path = _require_xlsx(path_text, settings)

    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"Excel-filen finns inte: {path_text}")

    try:
        max_rows = int(excel.get("max_rows_read", 100))
    except (TypeError, ValueError) as error:
        raise ValueError("excel.max_rows_read måste vara ett heltal.") from error

    if max_rows <= 0:
        raise ValueError("excel.max_rows_read måste vara större än 0.")

    openpyxl = openpyxl_module or _load_openpyxl()
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=False,
    )

    try:
        sheet = workbook.active
        rows = []

        for index, row in enumerate(
            sheet.iter_rows(values_only=True),
            start=1,
        ):
            if index > max_rows:
                break

            rows.append(list(row))

        return {
            "path": path,
            "sheet": sheet.title,
            "rows": rows,
            "truncated": sheet.max_row > max_rows,
            "total_rows": sheet.max_row,
            "total_columns": sheet.max_column,
        }
    finally:
        workbook.close()


def append_excel_row(
    request,
    settings=None,
    openpyxl_module=None,
):
    settings, excel = _excel_config(settings)

    if not excel.get("write_enabled", False):
        raise RuntimeError(
            "Excel-skrivning är avstängd i konfigurationen."
        )

    path = _require_xlsx(request["path"], settings)

    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"Excel-filen finns inte: {request['path']}")

    openpyxl = openpyxl_module or _load_openpyxl()
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active

    if sheet.max_column and len(request["row"]) != sheet.max_column:
        workbook.close()
        raise ValueError(
            "Den nya raden måste ha samma antal värden som arbetsbladet."
        )

    sheet.append(list(request["row"]))
    temporary = path.with_suffix(".tmp.xlsx")

    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()

        if temporary.exists():
            temporary.unlink()

    return {
        "path": path,
        "sheet": sheet.title,
        "row_number": sheet.max_row,
    }


def excel_create(user_text):
    try:
        request = parse_create_request(user_text)
        result = create_excel_file(request)
        return (
            f"Excel-filen {result['path'].name} skapades i workspace "
            f"med {result['columns']} kolumner och "
            f"{result['rows_written']} datarader."
        )
    except Exception as error:
        return f"Kunde inte skapa Excel-filen: {error}"


def excel_read(user_text):
    try:
        path = extract_file_path(user_text)

        if not path:
            return "Ingen .xlsx-fil hittades i instruktionen."

        result = read_excel_file(path)
        lines = [
            (
                f"Excel {result['path'].name} | blad {result['sheet']} | "
                f"{result['total_rows']} rader x "
                f"{result['total_columns']} kolumner"
            )
        ]

        for row in result["rows"]:
            lines.append(" | ".join(
                "" if value is None else str(value)
                for value in row
            ))

        if result["truncated"]:
            lines.append("[Visningen är trunkerad.]")

        return "\n".join(lines)
    except Exception as error:
        return f"Kunde inte läsa Excel-filen: {error}"


def excel_append(user_text):
    try:
        request = parse_append_request(user_text)
        result = append_excel_row(request)
        return (
            f"En rad lades till i {result['path'].name} "
            f"på rad {result['row_number']}."
        )
    except Exception as error:
        return f"Kunde inte lägga till Excel-raden: {error}"


TOOLS = {
    "excel_create": {
        "function": excel_create,
        "description": (
            "Skapar en .xlsx-tabell i workspace från kolumner och rader "
            "som anges i användarens instruktion."
        ),
        "input_mode": "user_text",
    },
    "excel_read": {
        "function": excel_read,
        "description": (
            "Läser värden och formler från en namngiven .xlsx-fil i workspace."
        ),
        "input_mode": "user_text",
    },
    "excel_append": {
        "function": excel_append,
        "description": (
            "Lägger till en rad i en befintlig .xlsx-fil i workspace."
        ),
        "input_mode": "user_text",
    },
}
